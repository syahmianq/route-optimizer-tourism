"""Console input and output for the tourism route optimiser."""

from itertools import permutations
from time import perf_counter

from openpyxl.styles import Alignment, Font

CAR_TYPES = {"Axia": 27.4, "Bezza": 21.1, "Myvi": 22.2, "Ativa": 18.9, "Alza": 22.0, "Aruz": 15.6, "Saga": 19.6, "Iriz": 13.5, "Persona": 12.5, "S70": 16.1}


def _show_error(*messages):
    width = max(len(message) for message in messages)
    print(f"\033[1;31m " + "_" * width + "______")
    print("|   " + " " * width + "   |")
    print("|   " + "ERROR!!!".center(width) + "   |")
    for message in messages:
        print("|   " + message.center(width) + "   |")
    print("|___" + "_" * width + "___|\033[0m")


def _choose_destinations(places):
    print("Welcome to your personal travel routing system.")
    print("Here are some of the popular places to visit in Kuala Lumpur:\n")
    for number, place in enumerate(places, start=1):
        print(f"   {number}. {place}")
    names_by_lowercase = {place.lower(): place for place in places}
    while True:
        user_input = input("\nEnter 3-5 destination numbers or names, separated by commas:\n")
        entries = [item.strip() for item in user_input.split(",") if item.strip()]
        selected = []
        invalid = []
        for entry in entries:
            if entry.isdigit() and 1 <= int(entry) <= len(places):
                selected.append(places[int(entry) - 1])
            elif entry.lower() in names_by_lowercase:
                selected.append(names_by_lowercase[entry.lower()])
            else:
                invalid.append(entry)
        if invalid:
            _show_error(f"Invalid selection(s): {invalid}.")
        elif not 3 <= len(selected) <= 5:
            _show_error(f"You must select between 3 to 5 destinations. You selected {len(selected)}.")
        else:
            return selected


def _write_selected_places(workbook, all_places_sheet, selected_sheet, selected_places, places, filename):
    selected_sheet.delete_rows(1, selected_sheet.max_row)
    # The all-places worksheet labels its matrix axes as N1, N2, etc.  Its
    # positions match the order of the application's attraction list.
    matrix_positions = {place.lower(): index for index, place in enumerate(places, start=2)}
    for index, place in enumerate(selected_places, start=2):
        selected_sheet.cell(row=1, column=index, value=place).alignment = Alignment(horizontal="center", vertical="center")
        selected_sheet.cell(row=index, column=1, value=place).alignment = Alignment(horizontal="center", vertical="center")
    for row, origin in enumerate(selected_places, start=2):
        for column, destination in enumerate(selected_places, start=2):
            source_row = matrix_positions[origin.lower()]
            source_column = matrix_positions[destination.lower()]
            selected_sheet.cell(row=row, column=column, value=all_places_sheet.cell(row=source_row, column=source_column).value).alignment = Alignment(horizontal="center", vertical="center")
    for cell in selected_sheet[1]:
        cell.font = Font(bold=True)
    for row in range(1, len(selected_places) + 2):
        selected_sheet.cell(row=row, column=1).font = Font(bold=True)
    for column in selected_sheet.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column)
        selected_sheet.column_dimensions[column[0].column_letter].width = max_length + 2
    workbook.save(filename)


def _find_route(selected_places, selected_sheet):
    """Exhaustively evaluate every visit order and return the exact best route."""
    distance_matrix = [
        [selected_sheet.cell(row=row, column=column).value for column in range(2, len(selected_places) + 2)]
        for row in range(2, len(selected_places) + 2)
    ]

    solver_started = perf_counter()
    best_order = None
    shortest_distance = float("inf")
    for order in permutations(range(len(selected_places))):
        route_distance = sum(
            distance_matrix[order[index]][order[index + 1]]
            for index in range(len(order) - 1)
        )
        if route_distance < shortest_distance:
            best_order = order
            shortest_distance = route_distance
    solver_seconds = perf_counter() - solver_started

    route = [selected_places[index] for index in best_order]
    return route, shortest_distance, solver_seconds


def _choose_car(total_distance):
    print("\n" + "=" * 80 + "\n\nAvailable car:")
    for car in CAR_TYPES:
        print("   - " + car)
    print("   - others")
    while True:
        car = input("\nWhat car are you driving? (from the list above): ").strip()
        if car.lower() == "others":
            try:
                consumption = float(input("Please state the fuel consumption of your car: "))
                if consumption <= 0:
                    raise ValueError
            except ValueError:
                _show_error("Fuel consumption must be a positive number.")
                continue
            return total_distance / consumption
        for name, consumption in CAR_TYPES.items():
            if car.lower() == name.lower():
                return total_distance / consumption
        _show_error(f"Invalid selection: {car}.", 'Choose "others" if your car is not listed.')


def _show_summary(route, total_distance, litres, solver_seconds):
    solver_minutes = solver_seconds / 60
    messages = ["You can follow this path to minimize total distance travel:", " --> ".join(route), f"Total Distance: {total_distance:.2f} km", f"Fuel cost: RM{litres * 1.99:.2f}", f"Total carbon emission: {litres * 2.3:.2f} kg", f"Model finish after {solver_minutes:.6f} minutes.", "I hope you have a safe journey. Thank you for using me!"]
    width = max(len(message) for message in messages)
    print("\n" + " " * 10 + "=" * (width + 10))
    print(" " * 10 + "||       " + "SUMMARY".center(width) + "       ||")
    print(" " * 10 + "=" * (width + 10))
    for message in messages:
        print(" " * 10 + "||   " + message.center(width) + "   ||")
    print(" " * 10 + "=" * (width + 10))


def run_application(workbook, all_places_sheet, selected_sheet, places, filename):
    """Run all terminal prompts and display the optimised trip summary."""
    selected_places = _choose_destinations(places)
    _write_selected_places(workbook, all_places_sheet, selected_sheet, selected_places, places, filename)
    route, total_distance, solver_seconds = _find_route(selected_places, selected_sheet)
    litres = _choose_car(total_distance)
    _show_summary(route, total_distance, litres, solver_seconds)
