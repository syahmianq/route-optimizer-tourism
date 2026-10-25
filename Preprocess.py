"""Prepare the attraction distance workbook used by the route optimiser."""

from pathlib import Path

import networkx as nx
import osmnx as ox
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.exceptions import InvalidFileException


PLACES = [
    "National Art Gallery", "Saloma Bridge", "KLCC", "KL Tower",
    "National Monument", "Merdeka Square", "Central Market", "Bukit Bintang",
    "Perdana Botanical Garden", "National Mosque", "Chinatown", "Royal Museum",
    "Little India", "Thean Hou Temple",
]

PLACE_COORDS = {
    "National Art Gallery": (3.1740359, 101.705126), "Saloma Bridge": (3.1609793, 101.7055954),
    "KLCC": (3.1574878, 101.7108034), "KL Tower": (3.1529983, 101.7016486),
    "National Monument": (3.1488546, 101.6838276), "Merdeka Square": (3.1490301, 101.6933323),
    "Central Market": (3.1450178, 101.6950955), "Bukit Bintang": (3.1467941, 101.7087459),
    "Perdana Botanical Garden": (3.1429252, 101.6825746), "National Mosque": (3.1419837, 101.6891824),
    "Chinatown": (3.144528, 101.6980254), "Royal Museum": (3.1324662, 101.6932807),
    "Little India": (3.1292845, 101.6815846), "Thean Hou Temple": (3.1219579, 101.6850929),
}


def _fit_column_widths(worksheet):
    for column in worksheet.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column)
        worksheet.column_dimensions[column[0].column_letter].width = max_length + 2


def _create_distance_workbook(filename):
    """Build and save a workbook containing the road-distance matrix."""
    workbook = Workbook()
    all_places_sheet = workbook.active
    all_places_sheet.title = "All_Places"

    for index, _ in enumerate(PLACES, start=1):
        node_name = f"N{index}"
        all_places_sheet.cell(row=1, column=index + 1, value=node_name).alignment = Alignment(horizontal="center", vertical="center")
        all_places_sheet.cell(row=index + 1, column=1, value=node_name).alignment = Alignment(horizontal="center", vertical="center")

    graph = ox.graph_from_point((3.18, 101.695), dist=6000, network_type="drive")
    matrix = [[0] * len(PLACES) for _ in PLACES]
    for row, origin in enumerate(PLACES):
        for column, destination in enumerate(PLACES):
            if row >= column:
                matrix[row][column] = matrix[column][row]
                continue
            lat1, lon1 = PLACE_COORDS[origin]
            lat2, lon2 = PLACE_COORDS[destination]
            node1 = ox.nearest_edges(graph, X=lon1, Y=lat1)[0]
            node2 = ox.nearest_edges(graph, X=lon2, Y=lat2)[0]
            try:
                matrix[row][column] = nx.shortest_path_length(graph, node1, node2, weight="length")
            except nx.NetworkXNoPath:
                matrix[row][column] = float("inf")

    for row in range(len(PLACES)):
        for column in range(len(PLACES)):
            all_places_sheet.cell(row=row + 2, column=column + 2, value=round(matrix[row][column] / 1000, 2)).alignment = Alignment(horizontal="center", vertical="center")

    for cell in all_places_sheet[1]:
        cell.font = Font(bold=True)
    for row in range(1, len(PLACES) + 2):
        all_places_sheet.cell(row=row, column=1).font = Font(bold=True)
    _fit_column_widths(all_places_sheet)
    workbook.save(filename)
    return workbook


def prepare_workbook(filename="KL_Attraction_Place.xlsx"):
    """Load the matrix or create it, then return a fresh selection sheet."""
    file_path = Path(filename)
    try:
        workbook = load_workbook(file_path) if file_path.exists() else _create_distance_workbook(file_path)
        all_places_sheet = workbook["All_Places"]
    except (InvalidFileException, KeyError):
        workbook = _create_distance_workbook(file_path)
        all_places_sheet = workbook["All_Places"]

    if "Selected_Destinations" in workbook.sheetnames:
        workbook.remove(workbook["Selected_Destinations"])
    selected_sheet = workbook.create_sheet("Selected_Destinations")
    return workbook, all_places_sheet, selected_sheet
