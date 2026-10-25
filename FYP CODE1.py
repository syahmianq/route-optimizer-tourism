from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.exceptions import InvalidFileException
from networkx.algorithms.approximation import traveling_salesman_problem
import os
import random
import osmnx as ox
import networkx as nx

filename = "KL_Attraction_Place.xlsx"

places =[
"National Art Gallery",
"Saloma Bridge",
"KLCC",
"KL Tower",
"National Monument",
"Merdeka Square",
"Central Market",
"Bukit Bintang",
"Perdana Botanical Garden",
"National Mosque",
"Chinatown",
"Royal Museum",
"Little India",
"Thean Hou Temple"
]

# Create or load workbook
if os.path.exists(filename): # If workbook exist, it runs this line
    try: # Run a block of code that might potentially raise an exception
        wb = load_workbook(filename)

        # If 2nd sheet exists, remove it 
        if "Selected_Destinations" in wb.sheetnames:
            wb.remove(wb["Selected_Destinations"])

        # Create a new sheet named "Selected_Destinations" and
        # assign as ws_selected
        ws_selected = wb.create_sheet("Selected_Destinations")

        # Assign 1st worksheet as ws_all
        ws_all = wb["All_Places"]

    except InvalidFileException:
        # Run this block of code if exception arise within try block
        wb = Workbook()
        ws_all = wb.active
        ws_all.title = "All_Places"
        ws_selected = wb.create_sheet("Selected_Destinations")
        
else: # Create new workbook 
    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "All_Places"
    ws_selected = wb.create_sheet("Selected_Destinations")

    # Fill matrix header (row + col names) 
    for i in range(len(places)): 
        place = f"N{i+1}"
        ws_all.cell(row=1, column=i+2, value=place).alignment = Alignment(horizontal="center", vertical="center")
        ws_all.cell(row=i+2, column=1, value=place).alignment = Alignment(horizontal="center", vertical="center")

    # Defining destination's coordinate
    place_coords = {
"National Art Gallery": (3.1740359, 101.705126),
"Saloma Bridge": (3.1609793, 101.7055954),
"KLCC": (3.1574878, 101.7108034),
"KL Tower": (3.1529983, 101.7016486),
"National Monument": (3.1488546, 101.6838276),
"Merdeka Square": (3.1490301,101.6933323),
"Central Market": (3.1450178, 101.6950955),
"Bukit Bintang": (3.1467941, 101.7087459),
"Perdana Botanical Garden": (3.1429252, 101.6825746),
"National Mosque": (3.1419837, 101.6891824),
"Chinatown": (3.144528, 101.6980254),
"Royal Museum": (3.1324662, 101.6932807),
"Little India": (3.1292845, 101.6815846),
"Thean Hou Temple": (3.1219579, 101.6850929)
    }

    # Build graph
    center_point = (3.18, 101.695)
    dist = 6000
    G = ox.graph_from_point(center_point, dist=dist, network_type="drive")

    # Prepare matrix
    n = len(places)
    dist_matrix = [[0]*n for _ in range(n)]

    for i in range(n):
        for j in range(n):

            if i == j:
                dist_matrix[i][j] = 0
                continue

            # use stored values (symmetry)
            if i > j:
                dist_matrix[i][j] = dist_matrix[j][i]
                continue

            # get coordinates
            lat1, lon1 = place_coords[places[i]]
            lat2, lon2 = place_coords[places[j]]

            edge1 = ox.nearest_edges(G, X=lon1, Y=lat1)
            edge2 = ox.nearest_edges(G, X=lon2, Y=lat2)

            node1 = edge1[0]
            node2 = edge2[0]

            try:
                meters = nx.shortest_path_length(G, node1, node2, weight="length")
            except nx.NetworkXNoPath:
                meters = float("inf")

            dist_matrix[i][j] = meters
            
    # Fill real OSMnx distances into Excel
    for i in range(2, len(places)+2):
        for j in range(2, len(places)+2):

            di = i - 2  # convert Excel row → matrix index
            dj = j - 2  # Excel col → matrix index

            distance_km = round(dist_matrix[di][dj] / 1000, 2)

            ws_all.cell(row=i, column=j, value=distance_km)
            ws_all.cell(row=i, column=j).alignment = Alignment(horizontal="center", vertical="center")

    # Style header 
    for cell in ws_all[1]:
        cell.font = Font(bold=True)
    for i in range(1, len(places)+2):
        ws_all.cell(row=i, column=1).font = Font(bold=True)

    # Auto-fit column width
    for col in ws_all.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws_all.column_dimensions[col[0].column_letter].width = max_length + 2

    wb.save(filename)

# Printing available places in the UI 
available_places = [ws_all.cell(row=i, column=1).value for i in range(2, ws_all.max_row + 1)]

print("Welcome to your personal travel routing system.")
print("Here are some of the popular places to visit in Kuala Lumpur:\n")
for p in places:
    print(" " * 3 + "- " + p)

# User input (with retry until valid)
available_lower = [p.lower() for p in places]

while True:
    user_input = input("\nPlease enter 3-5 destinations separated by commas (from the list above) : \n")
    selected = [x.strip() for x in user_input.split(",") if x.strip()]
    invalid = [x for x in selected if x.lower() not in available_lower]
    
    # Input validation and error handling
    if invalid: 
        msg1 = "ERROR!!!"
        msg2 = f"Invalid selection(s): {invalid}."
        width = len(msg2)
        print(f"\033[1;31m " + "_" * width + "______")
        print("|   " + " " * width + "   |")
        print("|   " + msg1.center(width) + "   |")
        print("|   " + msg2.center(width) + "   |")
        print("|___" + "_" * width + "___|\033[0m")
    elif not (3 <= len(selected) <= 5):
        msg1 = "ERROR!!!!"
        msg3 = f"You must select between 3 to 5 destinations. You already selected {len(selected)}."
        width = len(msg3)
        print(f"\033[1;31m " + "_" * width + "______")
        print("|   " + " " * width + "   |")
        print("|   " + msg1.center(width) + "   |")
        print("|   " + msg3.center(width) + "   |")
        print("|___" + "_" * width + "___|\033[0m")
    else:
        break  # Exit loop when all are valid and within limit

# Clear second sheet
ws_selected.delete_rows(1, ws_selected.max_row)

# Header
proper_selected = []
for i in range(len(selected)):
    # Match correct case from original list
    dest = selected[i]
    proper_name = places[available_lower.index(dest.lower())]
    proper_selected.append(proper_name)
    ws_selected.cell(row=1, column=i+2, value=proper_name).alignment = Alignment(horizontal="center", vertical="center")
    ws_selected.cell(row=i+2, column=1, value=proper_name).alignment = Alignment(horizontal="center", vertical="center")

for i in range(len(selected)):
    dest_row = selected[i]
    for j in range(len(selected)):
        dest_col = selected[j]

        # Find the same destinations in the first sheet
        r1 = available_lower.index(dest_row.lower()) + 2
        c1 = available_lower.index(dest_col.lower()) + 2

        # Copy the corresponding distance
        ws_selected.cell(row=i+2, column=j+2, value=ws_all.cell(row=r1, column=c1).value)
        ws_selected.cell(row=i+2, column=j+2).alignment = Alignment(horizontal="center", vertical="center")

# Style
for cell in ws_selected[1]:
    cell.font = Font(bold=True)
for i in range(1, len(selected)+2):
    ws_selected.cell(row=i, column=1).font = Font(bold=True)

# Auto-fit column width
for col in ws_selected.columns:
    max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    ws_selected.column_dimensions[col[0].column_letter].width = max_length + 2

wb.save(filename)

# Prepare edge for node-node connection to be input in the TSP
edges = []
for i in range(len(proper_selected)):
    for j in range(len(proper_selected)):
        if i < j:
            edges.append((proper_selected[i],proper_selected[j], ws_selected.cell(row = i+2,column = j+2).value))

F = nx.Graph()
F.add_weighted_edges_from(edges)

# Find a path visiting all nodes (TSP approximation)
tsp_path = traveling_salesman_problem(F, weight='weight',cycle= False)

# Calculate total distance along this path
total_distance = 0
for i in range(len(tsp_path)-1):
    total_distance += F[tsp_path[i]][tsp_path[i+1]]['weight']

# Getting what car the user drive
print("\n" + "=" * 80)
print("\nAvailable car:")
car_type = [
    ("Axia",27.4),
    ("Bezza",21.1),
    ("Myvi",22.2),
    ("Ativa",18.9),
    ("Alza",22.0),
    ("Aruz",15.6),
    ("Saga",19.6),
    ("Iriz",13.5),
    ("Persona",12.5),
    ("S70",16.1),
    ("others",)
    ]

for i in range(len(car_type)):
    print(" " * 3 + "-",car_type[i][0])

while True:
    car = input("\nWhat car are you driving? (from the list above) : ")
    
    # Selection for "other" menu
    if car == "others":
        print("\033[4m\nfuel consumption\tCategory\t\tTypical Car Type\033[0m")
        print("10-14.9km/L\t\tHigh Usage\t\tLarge SUV/off-road")
        print("15-19.9km/L\t\tMedium\t\t\tSedan/cross-over")
        print("20-24.9km/L\t\tEfficient\t\tCompact/hybrid")
        print("25-29.9km/L\t\tVery Efficient\t\tHybrid/eco\n")
        fuel = float(input("Please state the fuel consumption of your car (in number): "))
        litre = total_distance / fuel
        cost = litre * 1.99
        carbon_emission = litre * 2.3
        break
    
    found = False
    for i in range(len(car_type)):
        if car.lower() == car_type[i][0].lower():
            litre = total_distance / car_type[i][1]
            cost = litre * 1.99
            carbon_emission = litre * 2.3
            found = True
            break

    if found:
        break
    
    # Input validation and error handling
    else:
        msg1 = "ERROR!!!"
        msg2 = f"Invalid selection: {car}."
        msg3 = 'Please choose the option "Others", if it is not in the list.'
        width = max(len(msg2),len(msg3))
        print(f"\033[1;31m " + "_" * width + "______")
        print("|   " + " " * width + "   |")
        print("|   " + msg1.center(width) + "   |")
        print("|   " + msg2.center(width) + "   |")
        print("|   " + msg3.center(width) + "   |")
        print("|___" + "_" * width + "___|\033[0m")

title = "\033[1mSUMMARY\033[0m"
msg4 = "You can follow this path to minimize total distance travel:"
msg5 = " --> ".join(tsp_path)
msg6 = f"Total Distance: {total_distance:.2f} km"
msg7 = f"Fuel cost: RM{cost:.2f}"
msg8 = f"Total carbon emission: {carbon_emission:.2f} kg"
msg9 = "I hope you have a safe journey.Thank you for using me!"
box_width = max(len(msg4), len(msg5))

# Summary printing
print("\n")
print(" " * 10 + "====="+"=" * box_width + "=====")
print(" " * 10 + "||       " + title.center(box_width) + "       ||")
print(" " * 10 + "====="+"=" * box_width + "=====")
print(" " * 10 + "||   " + msg4.center(box_width) + "   ||")
print(" " * 10 + "||   " + msg5.center(box_width) + "   ||")
print(" " * 10 + "||   " + msg6.center(box_width) + "   ||")
print(" " * 10 + "||   " + msg7.center(box_width) + "   ||")
print(" " * 10 + "||   " + msg8.center(box_width) + "   ||")
print(" " * 10 + "||   " + msg9.center(box_width) + "   ||")
print(" " * 10 + "====="+"=" * box_width + "=====")